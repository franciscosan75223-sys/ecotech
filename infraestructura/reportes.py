from datetime import datetime
from pathlib import Path
from dominio.validaciones import validar_fecha

class Informes:
    def __init__(self, empleados, proyecto, departamento, registro_de_tiempo):
        self.empleados = empleados
        self.proyecto = proyecto
        self.departamento = departamento
        self.registro_de_tiempo = registro_de_tiempo

    def generar_informes(self):
        if isinstance(self.registro_de_tiempo, list):
            for nombre, encabezados, filas in self.obtener_secciones():
                print("\n" + nombre)
                if not filas:
                    print("Sin registros")
                for fila in filas:
                    for encabezado, valor in zip(encabezados, fila):
                        print(encabezado + ":", valor)
                    print()
                if nombre == "Tiempo":
                    columna_horas = encabezados.index("Horas")
                    print("Total de horas:", sum(fila[columna_horas] for fila in filas))
            return
        print("informe de tiempo")
        print("departamento", self.departamento)
        self.registro_de_tiempo.mostrar_registro()

    def obtener_secciones(self):
        if not isinstance(self.registro_de_tiempo, list):
            tiempo = self.registro_de_tiempo
            return [("Tiempo", ["Fecha", "Empleado", "Proyecto", "Horas", "Tareas", "Departamento"], [
                [validar_fecha(str(tiempo.fecha)), tiempo.empleado, tiempo.proyecto,
                 tiempo.get_horas_trabajadas(), tiempo.get_descripcion_de_tareas(), self.departamento]
            ])]

        empleados = []
        nombres_empleados = {}
        for empleado in self.empleados:
            nombres_empleados[empleado.get_id()] = empleado.get_nombre()
            empleados.append([
                empleado.get_id(), empleado.get_nombre(), empleado.get_direccion(),
                empleado.get_telefono(), empleado.get_correo(), empleado.get_fecha_inicio(),
                empleado.get_salario(), empleado.get_rol()
            ])

        proyectos = []
        nombres_proyectos = {}
        for proyecto in self.proyecto:
            nombres_proyectos[proyecto.get_id()] = proyecto.get_nombre()
            proyectos.append([
                proyecto.get_id(), proyecto.get_nombre(), proyecto.get_descripcion_proyecto(),
                proyecto.get_fecha_de_inicio(), proyecto.get_miembro(), proyecto.get_presupuesto()
            ])

        departamentos = []
        for departamento in self.departamento:
            departamentos.append([
                departamento.get_id(), departamento.get_nombre(), departamento.get_departamento(),
                departamento.get_gerente_id(),
                nombres_empleados.get(departamento.get_gerente_id(), "Sin vincular")
            ])

        tiempos = []
        for registro in self.registro_de_tiempo:
            tiempos.append([
                registro.get_id(), registro.get_fecha(), registro.get_id_empleado(),
                nombres_empleados.get(registro.get_id_empleado(), "Sin vincular"),
                registro.get_id_proyecto(),
                nombres_proyectos.get(registro.get_id_proyecto(), "Sin vincular"),
                registro.get_horas_trabajadas(), registro.get_tareas_realizadas()
            ])

        return [
            ("Empleados", ["ID", "Nombre", "Direccion", "Telefono", "Correo", "Inicio de contrato", "Salario", "Rol"], empleados),
            ("Proyectos", ["ID", "Nombre", "Descripcion", "Fecha de inicio", "Miembro", "Presupuesto"], proyectos),
            ("Departamentos", ["ID", "Nombre", "Descripcion", "ID del gerente", "Gerente"], departamentos),
            ("Tiempo", ["ID", "Fecha", "ID empleado", "Empleado", "ID proyecto", "Proyecto", "Horas", "Tareas"], tiempos)
        ]

    def preparar_ruta(self, extension, carpeta=None):
        if carpeta is None:
            # Sube un nivel para no guardar en infraestructura/reportes
            carpeta = Path(__file__).resolve().parent.parent / "reportes"
        carpeta = Path(carpeta).resolve()
        carpeta.mkdir(parents=True, exist_ok=True)
        momento = datetime.now().strftime("%Y-%m-%d_%H-%M-%S_%f")
        return carpeta / f"informe_ecotech_{momento}.{extension}"

    def exportar_excel(self, carpeta=None):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from math import ceil

        ruta = self.preparar_ruta("xlsx", carpeta)
        libro = Workbook()
        libro.remove(libro.active)

        for nombre, encabezados, filas in self.obtener_secciones():
            hoja = libro.create_sheet(nombre)
            hoja.append(encabezados)
            for fila in filas:
                hoja.append(fila)

            for fila in hoja:
                for celda in fila:
                    if isinstance(celda.value, str):
                        celda.data_type = "s"
                    if celda.is_date:
                        celda.number_format = "yyyy-mm-dd"
                    celda.alignment = Alignment(vertical="top", wrap_text=True)

            for celda in hoja[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="22624B")
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions

            for numero, encabezado in enumerate(encabezados, start=1):
                ancho = max([len(encabezado)] + [len(str(f[numero - 1] or "")) for f in filas])
                hoja.column_dimensions[get_column_letter(numero)].width = min(max(ancho + 3, 16), 48)
                if encabezado in ("Salario", "Presupuesto"):
                    for fila in hoja.iter_rows(min_row=2, min_col=numero, max_col=numero):
                        fila[0].number_format = '#,##0.00'
            for fila in hoja:
                lineas = 1
                for celda in fila:
                    ancho = hoja.column_dimensions[celda.column_letter].width - 3
                    partes = str(celda.value or "").splitlines() or [""]
                    lineas = max(lineas, sum(max(1, ceil(len(parte) / ancho)) for parte in partes))
                hoja.row_dimensions[fila[0].row].height = 16 * lineas + 6

            if not filas:
                hoja.append(["Sin registros"])
            if nombre == "Tiempo":
                columna_horas = encabezados.index("Horas")
                total = sum(fila[columna_horas] for fila in filas)
                hoja.append([])
                hoja.append(["Total de horas", total])
                hoja.cell(hoja.max_row, 1).font = Font(bold=True)

        libro.save(ruta)
        return ruta

    def exportar_pdf(self, carpeta=None):
        from html import escape
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

        ruta = self.preparar_ruta("pdf", carpeta)
        estilos = getSampleStyleSheet()
        estilos["BodyText"].spaceAfter = 4
        contenido = [Paragraph("EcoTech - Informe", estilos["Title"])]
        contenido.append(Paragraph("Generado: " + datetime.now().strftime("%Y-%m-%d %H:%M"), estilos["Normal"]))
        contenido.append(Spacer(1, 16))

        for numero_seccion, (nombre, encabezados, filas) in enumerate(self.obtener_secciones()):
            if numero_seccion > 0:
                contenido.append(PageBreak())
            contenido.append(Paragraph(nombre, estilos["Heading1"]))
            if not filas:
                contenido.append(Paragraph("Sin registros", estilos["BodyText"]))
            for numero_fila, fila in enumerate(filas, start=1):
                contenido.append(Paragraph(f"Registro {numero_fila}", estilos["Heading2"]))
                for encabezado, valor in zip(encabezados, fila):
                    if valor is None:
                        valor = "Sin vincular"
                    texto = escape(str(valor)).replace("\n", "<br/>")
                    contenido.append(Paragraph(f"<b>{escape(encabezado)}:</b> {texto}", estilos["BodyText"]))
                contenido.append(Spacer(1, 10))
            if nombre == "Tiempo":
                columna_horas = encabezados.index("Horas")
                total = sum(fila[columna_horas] for fila in filas)
                contenido.append(Paragraph(f"Total de horas: {total}", estilos["Heading2"]))

        documento = SimpleDocTemplate(str(ruta), pagesize=A4, title="Informe EcoTech",
                                      leftMargin=42, rightMargin=42, topMargin=42, bottomMargin=42)
        def poner_numero_pagina(lienzo, documento):
            lienzo.saveState()
            lienzo.setFont("Helvetica", 9)
            lienzo.drawRightString(A4[0] - 42, 24, f"Pagina {documento.page}")
            lienzo.restoreState()

        documento.build(contenido, onFirstPage=poner_numero_pagina, onLaterPages=poner_numero_pagina)
        return ruta