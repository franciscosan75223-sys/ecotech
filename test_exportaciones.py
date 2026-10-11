"""Comprueba archivos reales usando datos ficticios; no conecta con Supabase."""
import tempfile
import unittest
import io
from contextlib import redirect_stdout
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader
import main as app


def informe_de_prueba(cantidad=2):
    empleado = app.Empleado(1, 'José Muñoz', 'Calle de ejemplo 123', '912345678',
                            'jose@example.com', '2026-01-01', 1250, 'Gerente')
    proyecto = app.Proyecto(2, 'Energía limpia', 'Investigación & desarrollo',
                            '2026-01-01', 'Equipo de prueba', 2000)
    departamento = app.Departamento(3, 'Desarrollo', 1, 'Sostenibilidad')
    registros = []
    for numero in range(cantidad):
        horas = 8 if numero == 0 else 3
        tarea = f'Tarea {numero + 1}: revisión de energía & equipos <nuevos>. Segunda parte.'
        registros.append(app.Asignacion(numero + 1, '2026-01-10', horas, tarea, 1, 2))
    return app.Informes([empleado], [proyecto], [departamento], registros)


class PruebasExportaciones(unittest.TestCase):
    def setUp(self):
        self.carpeta = tempfile.TemporaryDirectory()
        self.addCleanup(self.carpeta.cleanup)

    def test_mostrar_informe_general(self):
        salida = io.StringIO()
        with redirect_stdout(salida):
            informe_de_prueba().generar_informes()
        self.assertIn('Total de horas: 11', salida.getvalue())
        self.assertEqual(salida.getvalue().count('Tarea 2:'), 1)

    def test_excel_contiene_las_cuatro_secciones_y_total(self):
        ruta = informe_de_prueba().exportar_excel(self.carpeta.name)
        libro = load_workbook(ruta)
        self.addCleanup(libro.close)
        self.assertEqual(libro.sheetnames, ['Empleados', 'Proyectos', 'Departamentos', 'Tiempo'])
        self.assertEqual(libro['Empleados']['B2'].value, 'José Muñoz')
        self.assertEqual(libro['Empleados']['G2'].value, 1250)
        self.assertEqual(libro['Empleados']['D2'].data_type, 's')
        self.assertIsInstance(libro['Tiempo']['B2'].value, datetime)
        self.assertEqual(libro['Tiempo']['G2'].value, 8)
        self.assertEqual(libro['Tiempo']['G3'].value, 3)
        self.assertEqual(libro['Tiempo']['A5'].value, 'Total de horas')
        self.assertEqual(libro['Tiempo']['B5'].value, 11)

    def test_excel_guarda_texto_sin_ejecutar_formulas(self):
        informe = informe_de_prueba(1)
        informe.registro_de_tiempo = [app.Asignacion(1, '2026-01-10', 8, '=SUM(1,1)', 1, 2)]
        ruta = informe.exportar_excel(self.carpeta.name)
        libro = load_workbook(ruta)
        self.addCleanup(libro.close)
        self.assertEqual(libro['Tiempo']['H2'].value, '=SUM(1,1)')
        self.assertEqual(libro['Tiempo']['H2'].data_type, 's')

    def test_pdf_preserva_todos_los_datos_y_acentos(self):
        ruta = informe_de_prueba().exportar_pdf(self.carpeta.name)
        texto = '\n'.join(p.extract_text() for p in PdfReader(ruta).pages)
        for esperado in ('José Muñoz', 'Energía limpia', 'Departamentos',
                         'equipos <nuevos>', 'Segunda parte.', 'Total de horas: 11'):
            self.assertIn(esperado, texto)
        self.assertEqual(texto.count('Tarea 1:'), 1)
        self.assertEqual(texto.count('Tarea 2:'), 1)

    def test_pdf_varias_paginas_conserva_ultimo_registro(self):
        ruta = informe_de_prueba(60).exportar_pdf(self.carpeta.name)
        paginas = PdfReader(ruta).pages
        self.assertGreater(len(paginas), 4)
        self.assertIn('Tarea 60:', paginas[-1].extract_text())
        self.assertIn('Total de horas: 185', paginas[-1].extract_text())

    def test_informe_vacio_y_nombres_distintos(self):
        informe = app.Informes([], [], [], [])
        pdf1 = informe.exportar_pdf(self.carpeta.name)
        pdf2 = informe.exportar_pdf(self.carpeta.name)
        self.assertNotEqual(pdf1, pdf2)
        self.assertTrue(pdf1.exists() and pdf2.exists())
        self.assertIn('Sin registros', PdfReader(pdf1).pages[0].extract_text())
        libro = load_workbook(informe.exportar_excel(self.carpeta.name))
        self.addCleanup(libro.close)
        self.assertEqual(libro['Tiempo']['A2'].value, 'Sin registros')
        self.assertEqual(libro['Tiempo']['B4'].value, 0)

    def test_ejemplo_original_de_un_solo_tiempo(self):
        tiempo = app.Tiempo('2026-01-10', 8, 'Tarea de ejemplo\nOtra línea', 'Juan Pérez', 'Proyecto A', False)
        informe = app.Informes(tiempo.empleado, tiempo.proyecto, 'Desarrollo', tiempo)
        libro = load_workbook(informe.exportar_excel(self.carpeta.name))
        self.addCleanup(libro.close)
        self.assertEqual(libro['Tiempo']['B2'].value, 'Juan Pérez')
        self.assertEqual(libro['Tiempo']['D2'].value, 8)
        texto = PdfReader(informe.exportar_pdf(self.carpeta.name)).pages[0].extract_text()
        self.assertIn('Desarrollo', texto)
        self.assertIn('Otra línea', texto)

    def test_menu_llama_a_las_dos_exportaciones(self):
        with patch('builtins.input', side_effect=['16', '17', '0']), \
             patch('builtins.print'), \
             patch.object(app, 'exportar_informe_desde_supabase') as exportar:
            app.menu(None, None)
        self.assertEqual([c.args for c in exportar.call_args_list], [(None, 'excel'), (None, 'pdf')])

    def test_error_al_guardar_permite_seguir_en_menu(self):
        with patch('builtins.input', side_effect=['16', '0']), patch('builtins.print') as imprimir, \
             patch.object(app, 'exportar_informe_desde_supabase', side_effect=PermissionError):
            app.menu(None, None)
        self.assertTrue(any('No se pudo guardar' in str(c) for c in imprimir.call_args_list))


if __name__ == '__main__':
    unittest.main()
